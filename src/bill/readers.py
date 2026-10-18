"""把上传的字节变成一张表。

只有本模块关心编码、分隔符、ZIP 成员和 Excel 格式；别的模块拿到的是已经
规整成「表头 + 数据行」的表。

账单文件跟普通 CSV 有个关键区别：前面有若干行账户说明，中间一行分隔符，
真正的表头在中间，末尾还有账户汇总。所以读取分成两步——先原样把行读出来
（``RawTable``），再由调用方拿着 profile 的列名去定位表头（``find_header``）。
"""

from __future__ import annotations

import csv
import io
import re
import zipfile
from collections import Counter
from dataclasses import dataclass, field
from datetime import date, datetime, time
from pathlib import PurePosixPath

from .config import MAX_BYTES, MAX_PDF_PAGES, MAX_ROWS
from .errors import UploadError

__all__ = [
    "ACCEPT_ATTR",
    "DELIMITED_SUFFIXES",
    "PDF_SUFFIXES",
    "SPREADSHEET_SUFFIXES",
    "Grid",
    "RawTable",
    "build_grid",
    "find_header",
    "read_table",
]

SPREADSHEET_SUFFIXES = {".xlsx", ".xlsm"}
DELIMITED_SUFFIXES = {".csv", ".tsv", ".txt"}
ZIP_SUFFIXES = {".zip"}
PDF_SUFFIXES = {".pdf"}

# 文件选择框里给出的类型，加上 .zip —— 支付宝的导出就是一个压缩包。
ACCEPT_ATTR = ",".join(
    sorted(SPREADSHEET_SUFFIXES | DELIMITED_SUFFIXES | ZIP_SUFFIXES | PDF_SUFFIXES)
)

_ZIP_MAGIC = b"PK\x03\x04"
_PDF_MAGIC = b"%PDF-"
_DELIMITERS = (",", "\t", ";", "|")

# UTF-8 之外的候选。注意不能照「第一个不报错的就用」来挑：cp1252 / latin-1 对
# 任何字节流都不抛异常，会悄悄把 GBK 账单解成乱码还不报错。所以下面要打分。
_LEGACY_ENCODINGS = ("gb18030", "big5", "cp1252", "latin-1")
_UTF8_BOM = b"\xef\xbb\xbf"


@dataclass
class RawTable:
    """一份文件解码后的原始行，还没决定哪一行是表头。"""

    filename: str
    kind: str  # "csv" | "xlsx" | "pdf"
    rows: list[list[str]]
    encoding: str = ""
    delimiter: str = ""
    truncated: bool = False
    member: str = ""  # ZIP 里的成员名，非压缩包为空
    notes: list[str] = field(default_factory=list)  # 给用户看的读取说明

    def describe(self) -> list[tuple[str, str]]:
        """预览页上那几块「文件信息」。"""
        facts = [("格式", self.kind.upper())]
        if self.member:
            facts.append(("压缩包成员", self.member))
        if self.encoding:
            facts.append(("编码", self.encoding))
        if self.delimiter:
            facts.append(("分隔符", {",": "逗号", "\t": "制表符", ";": "分号", "|": "竖线"}.get(self.delimiter, self.delimiter)))
        facts.append(("读到的行数", f"{len(self.rows):,}"))
        return facts


@dataclass
class Grid:
    """定位好表头的表：表头 + 数据行。"""

    header: list[str]
    rows: list[list[str]]
    header_index: int
    source: RawTable

    def column(self, name: str) -> int | None:
        """列名 -> 列号。找不到返回 None。"""
        try:
            return self.header.index(name)
        except ValueError:
            return None

    def cell(self, row: list[str], name: str) -> str:
        index = self.column(name)
        if index is None or index >= len(row):
            return ""
        return row[index].strip()


def read_table(
    filename: str,
    data: bytes,
    *,
    probes: tuple[str, ...] = (),
    encoding: str | None = None,
    delimiter: str | None = None,
) -> RawTable:
    """把上传的字节读成原始行。

    ``probes`` 是「这份文件里应该出现哪些列名」，同时用来给编码和 ZIP 成员打分。
    探测 profile 时传所有候选列名的并集。
    """
    name = (filename or "").strip()
    if not name:
        raise UploadError("请先选一个文件。")
    if not data:
        raise UploadError("这个文件是空的。")
    if len(data) > MAX_BYTES:
        raise UploadError(f"单个文件不能超过 {MAX_BYTES // (1024 * 1024)} MB。")

    suffix = PurePosixPath(name.replace("\\", "/")).suffix.lower()

    # 顺序很重要：.xlsx 本身就是一个 zip 包，所以必须先按后缀认 Excel，
    # 再考虑「靠魔数认出这是个压缩包」的情况（支付宝有时把 zip 导成 .csv）。
    if suffix in SPREADSHEET_SUFFIXES:
        return _read_excel(name, data)
    if suffix == ".xls":
        raise UploadError("不支持老的 .xls 格式——请另存为 .xlsx 或 .csv 再传。")
    if suffix in PDF_SUFFIXES or data[:5] == _PDF_MAGIC:
        return _read_pdf(name, data, probes)
    if suffix in ZIP_SUFFIXES or data[:4] == _ZIP_MAGIC:
        return _read_zip(name, data, probes, encoding, delimiter)
    if suffix in DELIMITED_SUFFIXES or not suffix:
        return _read_delimited(name, data, probes, encoding, delimiter)
    raise UploadError(f"不认识的格式「{suffix}」。支持 CSV、TSV、XLSX、PDF 和 ZIP。")


# --------------------------------------------------------------------------
# 分隔符文本（CSV / TSV）
# --------------------------------------------------------------------------


def _read_delimited(
    filename: str,
    data: bytes,
    probes: tuple[str, ...],
    encoding: str | None,
    delimiter: str | None,
) -> RawTable:
    text, used_encoding = _decode(data, [encoding] if encoding else [], probes)
    used_delimiter = _choose_delimiter(text, delimiter)

    rows: list[list[str]] = []
    truncated = False
    reader = csv.reader(io.StringIO(text, newline=""), delimiter=used_delimiter)
    for index, raw in enumerate(reader):
        if index >= MAX_ROWS:
            truncated = True
            break
        rows.append([cell.strip() for cell in raw])

    table = RawTable(
        filename=filename,
        kind="csv",
        rows=rows,
        encoding=used_encoding,
        delimiter=used_delimiter,
        truncated=truncated,
    )
    if truncated:
        table.notes.append(f"文件超过 {MAX_ROWS:,} 行，只读了前面这些。")
    return table


def _choose_delimiter(text: str, declared: str | None) -> str:
    """挑分隔符。

    不拿整段文本去跑 csv.Sniffer —— 账单前面的说明行里也带逗号和冒号，
    sniffer 很容易被带偏。这里改成看**结构一致性**：哪种分隔符切出来的列数
    在行间最稳定，就用哪种；一致程度相同再比列数。只切出单列的分隔符直接
    跳过，因为再怎么切也不像一张表。
    """
    if declared and declared != "auto":
        return "\t" if declared == "\\t" else declared

    sample = "\n".join(text.splitlines()[:50])
    if not sample.strip():
        return ","

    best, best_key = ",", (0.0, 0)
    for candidate in _DELIMITERS:
        widths = [
            len(row) for row in csv.reader(io.StringIO(sample), delimiter=candidate)
        ]
        if not widths:
            continue
        significant = [width for width in widths if width > 1]
        if not significant:
            continue
        width, occurrences = Counter(significant).most_common(1)[0]
        key = (occurrences / len(widths), width)
        if key > best_key:
            best, best_key = candidate, key
    return best


def _decode(data: bytes, declared: list[str], probes: tuple[str, ...]) -> tuple[str, str]:
    """解码，返回 (文本, 实际用的编码)。

    判定顺序是刻意排的：

    1. **合法的 UTF-8 就按 UTF-8 处理。** 这是很强的信号——真正的 GBK 字节流
       几乎不可能碰巧组成合法 UTF-8。这一步必须排在「按字符打分」**之前**：
       Monzo 的账单里有 emoji，那些字节被 gb18030 误解时会解出一堆汉字，于是
       在「中文占比」这一项上反而把正确的 UTF-8 比了下去，整份文件里所有
       非 ASCII 文字（含商家名里的重音字母）都会变乱码。
    2. 不是 UTF-8，再按 profile 声明的编码和通用候选**打分**，得分为
       「表头列名命中几个」+「中文字符占比」。这正是拦住 cp1252 把 GBK
       账单静默解成乱码的地方。同分时声明的编码优先。
    """
    if data[:3] == _UTF8_BOM:
        return data.decode("utf-8-sig"), "utf-8-sig"
    try:
        return data.decode("utf-8"), "utf-8"
    except UnicodeDecodeError:
        pass

    ordered: list[str] = []
    for candidate in (*declared, *_LEGACY_ENCODINGS):
        if candidate and candidate not in ("utf-8", "utf-8-sig") and candidate not in ordered:
            ordered.append(candidate)

    best: tuple[tuple[int, int, int], str, str] | None = None
    for rank, candidate in enumerate(ordered):
        try:
            text = data.decode(candidate)
        except (UnicodeDecodeError, LookupError):
            continue
        if "�" in text:
            continue
        hits = sum(1 for probe in probes if probe in text)
        # 先比表头命中，再比中文占比，最后才看声明的先后（rank 小的优先）。
        key = (hits, _cjk_count(text), -rank)
        if best is None or key > best[0]:
            best = (key, text, candidate)

    if best is None:
        raise UploadError("这个文件解不成文本，它真的是一份 CSV 吗？")
    return best[1], best[2]


def _cjk_count(text: str, limit: int = 4000) -> int:
    return sum(1 for char in text[:limit] if "一" <= char <= "鿿")


# --------------------------------------------------------------------------
# Excel
# --------------------------------------------------------------------------


def _read_excel(filename: str, data: bytes) -> RawTable:
    try:
        from openpyxl import load_workbook
        from openpyxl.utils.exceptions import InvalidFileException
    except ImportError:  # pragma: no cover - 依赖在 pyproject 里声明了
        raise UploadError("缺少 Excel 支持（openpyxl 没装）。")

    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except InvalidFileException:
        raise UploadError("这看起来不是一个有效的 Excel 文件。")
    except Exception as error:  # zipfile/XML 的报错类型五花八门
        raise UploadError(f"读不了这个工作簿：{error}")

    # 账单只会在一张表里，取第一张有内容的——不像以前的预览器那样把所有
    # sheet 拼起来，那对记账没有意义。
    rows: list[list[str]] = []
    sheet_name = ""
    truncated = False
    try:
        for worksheet in workbook.worksheets:
            collected: list[list[str]] = []
            for index, values in enumerate(worksheet.iter_rows(values_only=True)):
                if index >= MAX_ROWS:
                    truncated = True
                    break
                collected.append([_cell(value) for value in values])
            if any(cell for row in collected for cell in row):
                rows, sheet_name = collected, worksheet.title
                break
    finally:
        workbook.close()

    if not rows:
        raise UploadError("这个工作簿里没有内容。")

    table = RawTable(filename=filename, kind="xlsx", rows=rows, truncated=truncated)
    table.notes.append(f"工作表：{sheet_name}")
    if truncated:
        table.notes.append(f"文件超过 {MAX_ROWS:,} 行，只读了前面这些。")
    return table


def _cell(value: object) -> str:
    """把 Excel 单元格摊平成文本。"""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, datetime):
        # openpyxl 把纯日期的单元格给成午夜 datetime，按表格的显示习惯只留日期。
        if value.time() == time.min:
            return value.date().isoformat()
        return value.isoformat(sep=" ")
    if isinstance(value, (date, time)):
        return value.isoformat()
    return str(value).strip()


# --------------------------------------------------------------------------
# ZIP（支付宝的导出是个压缩包，里面装着明细 CSV 和汇总 CSV）
# --------------------------------------------------------------------------


def _read_zip(
    filename: str,
    data: bytes,
    probes: tuple[str, ...],
    encoding: str | None,
    delimiter: str | None,
) -> RawTable:
    try:
        archive = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile:
        raise UploadError("这个压缩包打不开，可能已经损坏。")

    members: list[RawTable] = []
    skipped: list[str] = []
    with archive:
        for info in archive.infolist():
            name = info.filename
            if info.is_dir():
                continue
            # macOS 打包时会塞进来的元数据副本。
            if name.startswith("__MACOSX/") or PurePosixPath(name).name.startswith("._"):
                skipped.append(name)
                continue
            member_name = PurePosixPath(name).name
            suffix = PurePosixPath(member_name).suffix.lower()
            if suffix not in SPREADSHEET_SUFFIXES | DELIMITED_SUFFIXES:
                skipped.append(name)
                continue
            try:
                payload = archive.read(info)
                if suffix in SPREADSHEET_SUFFIXES:
                    table = _read_excel(member_name, payload)
                else:
                    table = _read_delimited(member_name, payload, probes, encoding, delimiter)
            except UploadError:
                skipped.append(name)
                continue
            table.member = name
            members.append(table)

    if not members:
        raise UploadError("这个压缩包里没有能读的 CSV 或 Excel 文件。")

    # 谁能对上最多列名就用谁——支付宝包里明细表命中 11 个列名，汇总表一个都
    # 对不上，所以直接选明细表。并列时选行数多的那个。
    def score(table: RawTable) -> tuple[int, int]:
        text = "\n".join(",".join(row) for row in table.rows[:60])
        return (sum(1 for probe in probes if probe in text), len(table.rows))

    best = max(members, key=score)
    if skipped:
        best.notes.append(f"压缩包里跳过了 {len(skipped)} 个不相关的文件。")
    if len(members) > 1:
        best.notes.append(f"压缩包里有 {len(members)} 个表格，用了「{best.member}」。")
    best.filename = filename
    return best


# --------------------------------------------------------------------------
# PDF（文字版对账单）
# --------------------------------------------------------------------------


def _read_pdf(filename: str, data: bytes, probes: tuple[str, ...]) -> RawTable:
    """把文字版 PDF 的表格摊成「表头 + 数据行」。

    只做抽取，不做列映射：每页的表格交给 pdfplumber，重复的表头丢掉，单元格
    里的排版折行拼回去，剩下的跟 CSV 一样交给 ``find_header``。

    两个跟普通流水不同的地方：

    - **一页可能有多张表。** 信用卡合并账单把人民币和每种外币分成不同小节，
      一节一张表，页眉汇总表也长得像表。这里靠「第一行有没有命中配置里的
      列名」区分交易表和汇总表，而不是每页只挑一张。
    - **币种写在标题里，不在列里。** 比如「人民币交易明细」和
      「(GBP)外币交易明细」，一张表属于哪个小节由标题在页面上的纵向位置决定。
      识别到这种小节时，读取层给每行补一列「币别」，配置里把它映射到
      ``[columns].currency`` 就行。

    密码按中国银行导出的习惯，先试文件名去掉扩展名，再试完整文件名。改过名的
    文件解不开时明确说清楚——静默跳过会让用户以为文件是空的。
    """
    try:
        import pdfplumber
    except ImportError:
        raise UploadError("缺少 PDF 支持（pdfplumber 没装）。")

    name = PurePosixPath(filename.replace("\\", "/")).name
    stem = PurePosixPath(name).stem
    pdf = None
    last_kind = "broken"
    # 每次都用新的 BytesIO：失败的解析会把流读到一半，复用会报奇怪的错。
    for password in (None, stem, name):
        try:
            pdf = pdfplumber.open(io.BytesIO(data), password=password)
            break
        except Exception as error:  # pdfplumber 把 pdfminer 的异常包了一层
            last_kind = _pdf_error_kind(error)
            if last_kind != "password":
                break

    if pdf is None:
        if last_kind == "password":
            raise UploadError(
                "这个 PDF 有密码，用文件名（去掉 .pdf）解不开。"
                "中国银行流水导出的密码就是文件名去掉 .pdf，请确认文件名没被改过。"
            )
        if last_kind == "forbidden":
            raise UploadError("这个 PDF 不允许提取文字，请换一份能复制的导出版本。")
        if last_kind == "encrypted":
            raise UploadError("这个 PDF 的加密方式不支持，请用阅读器另存一份再传。")
        raise UploadError("这个 PDF 打不开，可能已经损坏。")

    rows: list[list[str]] = []
    header: list[str] | None = None
    base_width = 0
    add_currency_column = False
    current_currency = ""
    skipped_pages: list[int] = []
    truncated = False
    total_pages = 0

    try:
        total_pages = len(pdf.pages)
        pages = pdf.pages[:MAX_PDF_PAGES]
        # 先看整份文件有没有「××交易明细」小节。有才补币别列——普通流水本来就
        # 有币别列，多补一列会让同一份文件里出现两个「币别」。
        marks_by_page = [_pdf_currency_marks(page) for page in pages]
        has_sections = any(marks for marks in marks_by_page)

        for page_number, (page, marks) in enumerate(zip(pages, marks_by_page), 1):
            try:
                entries = _pdf_table_entries(page)
            except Exception:
                skipped_pages.append(page_number)
                continue

            used = False
            for top, table in entries:
                # 小节标题在表格上方，纵向位置最近的标题决定这张表的币种。
                for mark_top, code in marks:
                    if mark_top <= top:
                        current_currency = code

                cleaned = [
                    [_pdf_cell(cell) for cell in row] for row in table
                ]
                cleaned = [row for row in cleaned if any(cell for cell in row)]
                if not cleaned:
                    continue

                labels = [_pdf_header_label(cell) for cell in cleaned[0]]
                # 没有列名可参考时（直接调用读取层的场景），第一张表的第一行
                # 就按表头处理——跟 CSV 的默认约定一致；之后只认跟它一样的行。
                if probes:
                    looks_like_header = _header_hits(labels, probes) >= 2
                else:
                    looks_like_header = header is None or _clean_header(labels) == header
                if looks_like_header:
                    if header is None:
                        header = _clean_header(labels)
                        base_width = len(header)
                        add_currency_column = (
                            has_sections and _CURRENCY_COLUMN not in header
                        )
                        rows.append(
                            [*header, _CURRENCY_COLUMN]
                            if add_currency_column
                            else list(header)
                        )
                    # 后一个小节的重复表头只丢这一行，数据行照常留着。
                    rows.extend(
                        _pdf_emit(row, current_currency, add_currency_column)
                        for row in cleaned[1:]
                    )
                    used = True
                elif header is not None and _pdf_same_width(cleaned[0], base_width):
                    # 跨页续表没有表头，列数和规范表头一致就算它的数据。
                    rows.extend(
                        _pdf_emit(row, current_currency, add_currency_column)
                        for row in cleaned
                    )
                    used = True

                if len(rows) >= MAX_ROWS:
                    truncated = True
                    break

            # 小节标题可能出现在这一页最后一张表的下方（人民币那一节的表格
            # 从下一页才开始），也可能这一页根本没有交易表。把标题记下来，
            # 下一页继续用它，否则第二页开头那几行会丢掉币种。
            for _mark_top, code in marks:
                current_currency = code

            if not used:
                skipped_pages.append(page_number)
            if truncated:
                break
    finally:
        pdf.close()

    if header is None:
        raise UploadError(
            "这个 PDF 里没有找到表格。图片版（扫描件）需要先 OCR 成文字版再传。"
        )

    notes: list[str] = [f"共 {total_pages} 页。"]
    if total_pages > MAX_PDF_PAGES:
        notes.append(f"PDF 超过 {MAX_PDF_PAGES} 页，只读了前面这些。")
        truncated = True
    if skipped_pages:
        listed = "、".join(str(number) for number in skipped_pages[:5])
        notes.append(f"第 {listed} 页没有交易表，已跳过。")
    if truncated:
        notes.append(f"文件超过 {MAX_ROWS:,} 行，只读了前面这些。")

    return RawTable(
        filename=filename,
        kind="pdf",
        rows=rows[:MAX_ROWS],
        truncated=truncated,
        notes=notes,
    )


_CURRENCY_COLUMN = "币别"

# 合并账单的币种小节标题。外币那一节写成「(GBP)外币交易明细」，币种就在
# 括号里；人民币一节不写代码。顺序很重要：先认外币，它的标题里也带「交易明细」。
_SECTION_FCY = re.compile(r"[（(]\s*([A-Za-z]{3})\s*[)）]\s*外币\s*交易明细")
_SECTION_NAMED = (
    ("人民币", "CNY"),
    ("美元", "USD"),
    ("英镑", "GBP"),
    ("欧元", "EUR"),
    ("日元", "JPY"),
)


def _section_currency(text: str) -> str:
    """从一行标题里认出它开启的是哪个币种的小节。"""
    match = _SECTION_FCY.search(text)
    if match:
        return match.group(1).upper()
    if "交易明细" not in text:
        return ""
    for label, code in _SECTION_NAMED:
        if label in text:
            return code
    return ""


def _pdf_currency_marks(page) -> list[tuple[float, str]]:
    """这一页上每个币种小节标题的纵向位置。读不到文字就当作没有。"""
    try:
        lines = page.extract_text_lines() or []
    except Exception:
        return []
    marks: list[tuple[float, str]] = []
    for line in lines:
        code = _section_currency(str(line.get("text") or ""))
        if code:
            marks.append((float(line.get("top") or 0), code))
    return marks


def _pdf_table_entries(page) -> list[tuple[float, list[list[str | None]]]]:
    """一页里的所有表格，按从上到下的顺序。"""
    entries: list[tuple[float, list[list[str | None]]]] = []
    for found in page.find_tables() or []:
        try:
            table = found.extract()
        except Exception:
            continue
        if table:
            entries.append((float(found.bbox[1]), table))
    entries.sort(key=lambda item: item[0])
    return entries


def _pdf_emit(
    row: list[str], currency: str, add_currency_column: bool
) -> list[str]:
    return [*row, currency] if add_currency_column else list(row)


def _pdf_same_width(row: list[str], width: int) -> bool:
    return len(row) == width


def _header_hits(labels: list[str], probes: tuple[str, ...]) -> int:
    """这一行里有多少格正好是配置认得的列名。"""
    wanted = {probe.strip() for probe in probes if probe.strip()}
    return len(wanted & {label.strip() for label in labels})


def _pdf_header_label(cell: str) -> str:
    """双语表头只留中文部分。

    中行的表头是一个格子里「中文\\nEnglish」，拼回一行后变成
    「交易日Transaction Date」，跟配置里的「交易日」对不上。这里在第一个
    ASCII 字母前截断——只有截出来的部分含中文才截，避免误伤纯英文表头。
    """
    text = cell.strip()
    if not any("一" <= char <= "鿿" for char in text):
        return text
    for index, char in enumerate(text):
        if char in _ASCII_LETTERS:
            prefix = text[:index].strip()
            return prefix or text
    return text


_ASCII_LETTERS = set("abcdefghijklmnopqrstuvwxyzABCDEFGHIJKLMNOPQRSTUVWXYZ")


def _pdf_cell(value: object) -> str:
    """把一个 PDF 单元格规整成一行文字。

    单元格里的换行是排版折行，不是内容：中文、数字、括号两侧直接相连
    （``万科支\\n行`` → ``万科支行``），只有两个 ASCII 字母之间才补空格
    （``FLYWIRE\\nPAYMENTS`` → ``FLYWIRE PAYMENTS``）。
    """
    text = "" if value is None else str(value)
    text = text.replace("\r", "").replace("\xa0", " ")
    parts = text.split("\n")
    joined = parts[0]
    for part in parts[1:]:
        if joined[-1:] in _ASCII_LETTERS and part[:1] in _ASCII_LETTERS:
            joined += " "
        joined += part
    return " ".join(joined.split()).strip()


def _pdf_error_kind(error: BaseException | None) -> str:
    """沿异常链找 pdfminer 的真实原因。

    pdfplumber 把 pdfminer 的异常包成 ``PdfminerException``，``str()`` 是空串；
    真正的原因在 ``args[0]`` 或 ``__context__`` 里，所以不能靠错误文字判断。
    """
    seen: set[int] = set()
    stack: list[object] = [error]
    while stack:
        current = stack.pop()
        if current is None or id(current) in seen:
            continue
        seen.add(id(current))
        name = type(current).__name__
        if name == "PDFPasswordIncorrect":
            return "password"
        if name == "PDFTextExtractionNotAllowed":
            return "forbidden"
        if name == "PDFEncryptionError":
            return "encrypted"
        stack.extend(getattr(current, "args", ()))
        stack.append(getattr(current, "__cause__", None))
        stack.append(getattr(current, "__context__", None))
    return "broken"


# --------------------------------------------------------------------------
# 表头定位
# --------------------------------------------------------------------------


def find_header(
    rows: list[list[str]],
    candidates: tuple[str, ...],
    *,
    scan: int = 60,
    start_after: str | None = None,
) -> int | None:
    """在开头若干行里找表头行，返回行号；找不到返回 None。

    这里刻意不做「挑最宽的那一行」之类的猜测——猜错表头比明确报错更糟。
    找不到就老实返回 None，让上层提示用户换一个配置。
    """
    if not candidates:
        return None

    start = 0
    if start_after:
        for index, row in enumerate(rows[:scan]):
            if any(start_after in cell for cell in row):
                start = index + 1
                break

    stop = min(len(rows), start + scan)
    for index in range(start, stop):
        cells = {cell.strip() for cell in rows[index]}
        if all(candidate in cells for candidate in candidates):
            return index
    return None


def build_grid(source: RawTable, header_index: int) -> Grid:
    """把原始行切成表头和数据行。"""
    header = _clean_header(source.rows[header_index])
    width = len(header)
    body = [(row + [""] * width)[:width] for row in source.rows[header_index + 1 :]]
    return Grid(header=header, rows=body, header_index=header_index, source=source)


def _clean_header(header: list[str]) -> list[str]:
    """补齐空列名、给重名的列名加序号，保持原有顺序。"""
    labels: list[str] = []
    seen: dict[str, int] = {}
    for index, raw in enumerate(header):
        label = raw.strip().replace("　", " ") if raw else ""
        label = label or f"第 {index + 1} 列"
        if label in seen:
            seen[label] += 1
            label = f"{label} ({seen[label]})"
        else:
            seen[label] = 0
        labels.append(label)
    return labels
